"""
dealbar: rendert alle Zeilen mit Status "neu" aus deals.xlsx zu Videos.
Wird von START.bat aufgerufen.
"""

import os
import re
import sys
import traceback
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from openpyxl import load_workbook  # noqa: E402

import brand as brandlib  # noqa: E402
from deal_video import render, euro  # noqa: E402

TABLE = os.path.join(ROOT, "deals.xlsx")
DEALS_DIR = os.path.join(ROOT, "deals")
VIDEO_DIR = os.path.join(ROOT, "videos")
IMG_EXT = (".jpg", ".jpeg", ".png", ".webp")

COLS = ["Status", "Ordner", "Produktname", "Alter Preis (€)", "Neuer Preis (€)", "Rabatt",
        "Preis geprüft am", "Hook", "Affiliate-Link", "Label", "Preis zeigen", "Erstellt am", "Videodatei"]


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("€", "").strip()
    if "," in s:  # deutsche Schreibweise 1.299,99
        s = s.replace(".", "").replace(",", ".")
    return float(s)


def stamp_text(v):
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y, %H:%M")
    if v:
        return str(v).replace(" Uhr", "").strip()
    return datetime.now().strftime("%d.%m.%Y, %H:%M")


def slug(s):
    s = s.lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:40] or "deal"


def caption(cfg, link, discount, stamp):
    lines = [cfg["title"]]
    if cfg.get("hook"):
        lines.append(cfg["hook"])
    lines.append("")
    if cfg.get("show_price", True) and cfg.get("price_new") is not None:
        p = f"Jetzt {euro(cfg['price_new'])}"
        if cfg.get("price_old"):
            p += f" statt {euro(cfg['price_old'])}"
            if discount:
                p += f" (-{discount} %)"
        lines.append(p + "*")
    lines.append("👉 Link in Bio")
    lines.append("")
    if cfg.get("show_price", True):
        lines.append(f"*Preis vom {stamp} Uhr, kann sich jederzeit ändern.")
    lines.append("#Werbung | Als Amazon-Partner verdiene ich an qualifizierten Verkäufen.")
    lines.append("#dealbar #amazondeals #angebot #schnäppchen #deal")
    if link:
        lines += ["", "-----", "Affiliate-Link (für deine Link-in-Bio-Seite):", link]
    return "\n".join(lines) + "\n"


def main():
    print("dealbar Video-Generator\n")
    if not os.path.exists(TABLE):
        sys.exit(f"Tabelle nicht gefunden: {TABLE}")
    # Pruefen, ob die Tabelle gespeichert werden kann (Excel geoeffnet?)
    try:
        with open(TABLE, "r+b"):
            pass
    except PermissionError:
        sys.exit("deals.xlsx ist noch in Excel geöffnet. Bitte Excel schließen und erneut starten.")

    brand = brandlib.load_brand(os.path.join(ROOT, "marke.json"))
    wb = load_workbook(TABLE)
    ws = wb["Deals"]
    head = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    missing = [c for c in COLS if c not in head]
    if missing:
        sys.exit("In deals.xlsx fehlen Spalten: " + ", ".join(missing))

    def cell(row, name):
        return ws.cell(row=row, column=head[name])

    todo = [r for r in range(2, ws.max_row + 1)
            if str(cell(r, "Status").value or "").strip().lower() == "neu" and cell(r, "Produktname").value]
    if not todo:
        print('Keine Zeilen mit Status "neu" gefunden. Nichts zu tun.')
        return

    os.makedirs(VIDEO_DIR, exist_ok=True)
    ok = 0
    for n, r in enumerate(todo, 1):
        title = str(cell(r, "Produktname").value).strip()
        print(f"[{n}/{len(todo)}] {title[:60]}")
        try:
            folder = str(cell(r, "Ordner").value or "").strip()
            if not folder:
                raise ValueError("Spalte 'Ordner' ist leer.")
            fpath = os.path.join(DEALS_DIR, folder)
            if not os.path.isdir(fpath):
                raise ValueError(f"Ordner nicht gefunden: deals\\{folder}")
            imgs = sorted(f for f in os.listdir(fpath) if f.lower().endswith(IMG_EXT))
            if not imgs:
                raise ValueError(f"Keine Bilder (jpg/png/webp) in deals\\{folder}")

            p_old, p_new = num(cell(r, "Alter Preis (€)").value), num(cell(r, "Neuer Preis (€)").value)
            stamp = stamp_text(cell(r, "Preis geprüft am").value)
            show_price = str(cell(r, "Preis zeigen").value or "ja").strip().lower() not in ("nein", "n", "no")
            discount = round((p_old - p_new) / p_old * 100) if (p_old and p_new and p_old > p_new) else None

            cfg = {
                "title": title,
                "images": [os.path.join(folder, f) for f in imgs[:4]],
                "price_old": p_old if (p_old and p_new and p_old > p_new) else None,
                "price_new": p_new,
                "price_time": stamp,
                "hook": str(cell(r, "Hook").value or "").strip(),
                "label": str(cell(r, "Label").value or "TOP-DEAL").strip(),
                "show_price": show_price,
                "show_discount": show_price and bool(discount),
            }
            if discount:
                cfg["discount_percent"] = discount

            name = f"{datetime.now():%Y-%m-%d}_{slug(folder)}"
            video = os.path.join(VIDEO_DIR, name + ".mp4")
            render(cfg, DEALS_DIR, video, brand)
            link = str(cell(r, "Affiliate-Link").value or "").strip()
            with open(os.path.join(VIDEO_DIR, name + "_beschreibung.txt"), "w", encoding="utf-8") as fh:
                fh.write(caption(cfg, link, discount, stamp))

            cell(r, "Status").value = "fertig"
            cell(r, "Erstellt am").value = datetime.now().strftime("%d.%m.%Y %H:%M")
            cell(r, "Videodatei").value = name + ".mp4"
            ok += 1
            print(f"      fertig: videos\\{name}.mp4")
        except Exception as e:
            cell(r, "Status").value = "Fehler"
            cell(r, "Videodatei").value = str(e)[:200]
            print(f"      FEHLER: {e}")
            if os.environ.get("DEALBAR_DEBUG"):
                traceback.print_exc()

    try:
        wb.save(TABLE)
    except PermissionError:
        print("\nHinweis: deals.xlsx konnte nicht aktualisiert werden (in Excel geöffnet?).")
    print(f"\n{ok} von {len(todo)} Videos erstellt. Sie liegen im Ordner 'videos'.")


if __name__ == "__main__":
    main()
