"""
dealbar: gemeinsame Bausteine fuer Telegram-Bot, Video-Generator und Website.

- deals.xlsx lesen und schreiben, mit Sperre (Bot und Video-Generator laufen gleichzeitig)
  und ohne halb geschriebene Dateien (erst in eine Hilfsdatei speichern, dann austauschen)
- fehlende Spalten ("Nr", "ASIN") werden beim ersten Schreiben automatisch ergaenzt
- Tabellen-Konflikte zeilenweise zusammenfuehren (falls deals.xlsx auch auf GitHub geaendert wurde)
- Amazon-Links aufraeumen und mit deiner Partner-ID versehen
- Einstellungen aus website.json
"""

import contextlib
import copy
import io
import json
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TABLE = os.path.join(ROOT, "deals.xlsx")
DEALS_DIR = os.path.join(ROOT, "deals")
VIDEO_DIR = os.path.join(ROOT, "videos")
SETTINGS = os.path.join(ROOT, "website.json")
LOCK = os.path.join(ROOT, ".deals.lock")

IMG_EXT = (".jpg", ".jpeg", ".png", ".webp")
SCREENSHOT = ("screenshot", "bildschirmfoto")
HIDDEN = ("abgelaufen", "aus", "offline")
DISCLOSURE = "Als Amazon-Partner verdiene ich an qualifizierten Verkäufen."

# Kategorie -> (Emoji, kurzer Name fuer Knoepfe, Hashtag)
CATEGORIES = {
    "Technik": ("💻", "Technik", "technik"),
    "Gaming": ("🎮", "Gaming", "gaming"),
    "Haushalt": ("🏠", "Haushalt", "haushalt"),
    "Küche": ("🍳", "Küche", "küche"),
    "Mode": ("👕", "Mode", "mode"),
    "Beauty & Pflege": ("💄", "Beauty", "beauty"),
    "Sport & Freizeit": ("⚽", "Sport", "sport"),
    "Spielzeug": ("🧸", "Spielzeug", "spielzeug"),
    "Baby & Kind": ("🍼", "Baby", "baby"),
    "Garten & Baumarkt": ("🔨", "Garten", "garten"),
    "Tierbedarf": ("🐾", "Tiere", "haustier"),
    "Sonstiges": ("📦", "Sonstiges", "angebot"),
}
ALIASES = {
    "kleidung": "Mode", "fashion": "Mode", "bekleidung": "Mode", "schuhe": "Mode", "jacke": "Mode",
    "schuhe & handtaschen": "Mode", "elektronik": "Technik", "tech": "Technik", "computer": "Technik",
    "elektronik & foto": "Technik", "computer & zubehör": "Technik", "handy": "Technik", "audio": "Technik",
    "smartphone": "Technik", "games": "Gaming", "konsole": "Gaming", "videospiele": "Gaming",
    "küche, haushalt & wohnen": "Haushalt", "wohnen": "Haushalt", "deko": "Haushalt", "putzen": "Haushalt",
    "kueche": "Küche", "kochen": "Küche", "kaffee": "Küche", "drogerie": "Beauty & Pflege",
    "pflege": "Beauty & Pflege", "kosmetik": "Beauty & Pflege", "drogerie & körperpflege": "Beauty & Pflege",
    "beauty": "Beauty & Pflege", "sport": "Sport & Freizeit", "fitness": "Sport & Freizeit",
    "outdoor": "Sport & Freizeit", "freizeit": "Sport & Freizeit", "fahrrad": "Sport & Freizeit",
    "spiele": "Spielzeug", "lego": "Spielzeug", "baby": "Baby & Kind", "kind": "Baby & Kind",
    "kinder": "Baby & Kind", "garten": "Garten & Baumarkt", "werkzeug": "Garten & Baumarkt",
    "baumarkt": "Garten & Baumarkt", "heimwerken": "Garten & Baumarkt", "tiere": "Tierbedarf",
    "tier": "Tierbedarf", "haustier": "Tierbedarf", "hund": "Tierbedarf", "katze": "Tierbedarf",
    "sonstige": "Sonstiges",
}

STATUS_LIST = "neu,fertig,Fehler,abgelaufen"
NEW_COLUMNS = (("Nr", 7), ("ASIN", 14))


# --------------------------------------------------------------------------- kleine Helfer

def num(v):
    """Zahl aus Zelle oder Text: 19,99 | 19.99 | 1.299,00 | 1,299.00 | 19 € -> float (sonst None)."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not re.search(r"\d", s):
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?[1-9]\d{0,2}(\.\d{3})+", s):
        s = s.replace(".", "")  # 1.299 = eintausendzweihundertneunundneunzig
    try:
        return float(s)
    except ValueError:
        return None


def euro(v):
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{s} €"


def discount(old, new):
    return round((old - new) / old * 100) if (old and new and old > new) else None


def slug(s):
    s = str(s).lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:40].strip("-") or "deal"


def parse_dt(v):
    """'25.10.2026, 08:59' | '25.10.2026 08:59 Uhr' | '2026-10-25T08:59' | datetime -> datetime (sonst None)."""
    if isinstance(v, datetime):
        return v
    s = str(v or "")
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})(?:\D{1,5}(\d{1,2}):(\d{2}))?", s)
    if m:
        d, mo, y, h, mi = m.groups()
    else:
        m = re.search(r"(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2}))?", s)
        if not m:
            return None
        y, mo, d, h, mi = m.groups()
    try:
        return datetime(int(y), int(mo), int(d), int(h or 0), int(mi or 0))
    except ValueError:
        return None


def fmt_dt(dt):
    return dt.strftime("%d.%m.%Y, %H:%M") if dt else ""


def iso(dt):
    """ISO-Zeit mit Zeitzone (fuer den Browser), z. B. 2026-10-25T08:59:00+01:00.
    Zeiten in der Tabelle sind deutsche Ortszeit, egal auf welchem Rechner gebaut wird."""
    if not dt:
        return ""
    if dt.tzinfo is None:
        try:
            from zoneinfo import ZoneInfo
            dt = dt.replace(tzinfo=ZoneInfo("Europe/Berlin"))
        except Exception:  # Windows ohne Zeitzonen-Daten: Zeitzone des PCs
            dt = dt.astimezone()
    return dt.isoformat(timespec="seconds")


def match_category(text):
    """Ganze Zeile -> Kategorie (oder None). 'Kleidung' -> 'Mode', 'Fashion' -> 'Mode' ..."""
    low = re.sub(r"\s+", " ", str(text or "").lower()).strip(" .:#!")
    if not low or len(low) > 40:
        return None
    if low in ALIASES:
        return ALIASES[low]
    for c, (_, short, _) in CATEGORIES.items():
        if low in (c.lower(), short.lower()) or low in [p.strip() for p in c.lower().split("&")]:
            return c
    return None


def images_in(folder):
    """Eigene Bilder eines Deal-Ordners (sortiert, ohne Screenshots)."""
    path = os.path.join(DEALS_DIR, folder or "")
    if not folder or not os.path.isdir(path):
        return []
    return [f for f in sorted(os.listdir(path))
            if f.lower().endswith(IMG_EXT) and not f.lower().startswith(SCREENSHOT)]


# --------------------------------------------------------------------------- Amazon-Links

URL_RE = re.compile(r"(?:https?://|(?<![\w./])(?=(?:www\.)?(?:amazon\.[a-z.]{2,6}|amzn\.(?:to|eu))/))[^\s<>\"'«»]+",
                    re.I)
SHORT_HOSTS = ("amzn.to", "amzn.eu", "a.co", "amzn.com")
ASIN_RE = re.compile(r"/(?:dp|gp/product|gp/aw/d|gp/offer-listing|exec/obidos/asin|o/asin|product-reviews)/"
                     r"([A-Z0-9]{10})(?=[/?&#]|$)", re.I)
UA = {"User-Agent": "Mozilla/5.0 (Linux; Android 14) dealbar-bot"}


def host_of(url):
    try:
        return (urllib.parse.urlparse(url).hostname or "").lower()
    except ValueError:
        return ""


def is_amazon(url):
    h = re.sub(r"^(www|smile|m)\.", "", host_of(url))
    return h in SHORT_HOSTS or bool(re.fullmatch(r"amazon\.(de|at|com|co\.uk|fr|it|es|nl|pl|se|com\.be|com\.tr)", h))


def find_link(text):
    """Ersten Amazon-Link im Text finden -> (link, start, ende) oder None."""
    for m in URL_RE.finditer(text or ""):
        url = m.group(0).rstrip(").,;:!?\"'")
        if not url.lower().startswith("http"):
            url = "https://" + url
        if is_amazon(url):
            return url, m.start(), m.start() + len(m.group(0))
    return None


def asin_of(url):
    h = host_of(url)
    if not h or re.sub(r"^(www|smile|m)\.", "", h) in SHORT_HOSTS or "amazon." not in h:
        return None
    p = urllib.parse.urlparse(url)
    m = ASIN_RE.search(p.path)
    if m:
        return m.group(1).upper()
    q = urllib.parse.parse_qs(p.query)
    for k in ("asin", "ASIN"):
        if q.get(k) and re.fullmatch(r"[A-Z0-9]{10}", q[k][0], re.I):
            return q[k][0].upper()
    return None


class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs):
        return None


def resolve_short(url, hops=4):
    """Kurzlink (amzn.eu/d/..., amzn.to/...) aufloesen, ohne Amazon-Seiten selbst abzurufen."""
    opener = urllib.request.build_opener(_NoRedirect)
    for _ in range(hops):
        if re.sub(r"^www\.", "", host_of(url)) not in SHORT_HOSTS:
            return url
        loc = None
        for method in ("HEAD", "GET"):
            try:
                with opener.open(urllib.request.Request(url, method=method, headers=UA), timeout=15) as r:
                    loc = r.headers.get("Location")
            except urllib.error.HTTPError as e:
                loc = e.headers.get("Location")
            except Exception as e:  # Netzwerkfehler: Link bleibt, wie er ist
                print("Kurzlink nicht aufgeloest:", e)
                return url
            if loc:
                break
        if not loc:
            return url
        url = urllib.parse.urljoin(url, loc)
    return url


def amazon_domain(url):
    h = re.sub(r"^(www|smile|m)\.", "", host_of(url))
    return "www." + h if h.startswith("amazon.") else "www.amazon.de"


def clean_link(url, tag="", resolve=True):
    """Amazon-Link aufraeumen -> (Partnerlink, ASIN, Hinweis oder None).

    Aus jeder Produkt-Adresse wird https://www.amazon.de/dp/<ASIN>?tag=<deine Partner-ID>.
    Fremde Partner-IDs (z. B. aus weitergeleiteten Deals anderer Kanaele) fliegen immer raus.
    """
    final = url
    asin = asin_of(url)
    if not asin and resolve and re.sub(r"^www\.", "", host_of(url)) in SHORT_HOSTS:
        final = resolve_short(url)
        asin = asin_of(final)
    if not asin:
        return url, None, "Produktnummer (ASIN) nicht erkannt. Am besten den Link aus der Amazon-App teilen."
    domain = amazon_domain(final)
    link = f"https://{domain}/dp/{asin}" + (f"?tag={tag}" if tag else "")
    note = None if domain == "www.amazon.de" else "Achtung: Der Link ist nicht von amazon.de, die Partner-ID gilt nur dort."
    return link, asin, note


def with_tag(link, tag):
    """Partner-ID in einen gespeicherten Link einsetzen (fuer Deals von vor der Partner-ID)."""
    asin = asin_of(link or "")
    if not (tag and asin) or amazon_domain(link) != "www.amazon.de":
        return link
    return f"https://www.amazon.de/dp/{asin}?tag={tag}"


def plain_link(link):
    """Produktlink ohne Partner-ID (zum Pruefen im eigenen Chat, damit keine Eigenkaeufe ueber die ID laufen)."""
    asin = asin_of(link or "")
    return f"https://{amazon_domain(link)}/dp/{asin}" if asin else link


def valid_link(link):
    link = str(link or "").strip()
    return bool(re.match(r"https?://", link)) and "DEIN-LINK" not in link


# --------------------------------------------------------------------------- Einstellungen

DEFAULT_SETTINGS = {
    "titel": "dealbar – geprüfte Amazon-Deals, täglich frisch",
    "beschreibung": "Die besten Amazon-Angebote des Tages: von Hand geprüft, mit Rabatt auf einen Blick.",
    "domain": "",
    "rabatt_zeigen_stunden": 24,
    "preise_in_euro_zeigen": False,
    "aeltere_deals_ab_tagen": 21,
    "partner_tag": "",
    "telegram_kanal": "",
    "social": {"instagram": "", "tiktok": "", "youtube": "", "pinterest": "", "facebook": ""},
    "impressum": {"name": "VORNAME NACHNAME", "strasse": "STRASSE HAUSNUMMER", "plz_ort": "PLZ ORT",
                  "email": "EMAIL@BEISPIEL.DE", "telefon": ""},
}
PLACEHOLDER = re.compile(r"VORNAME|NACHNAME|STRASSE|PLZ ORT|BEISPIEL\.DE")


def load_settings():
    s = copy.deepcopy(DEFAULT_SETTINGS)
    if os.path.exists(SETTINGS):
        with open(SETTINGS, encoding="utf-8-sig") as fh:
            data = json.load(fh)
        for k, v in data.items():
            if isinstance(v, dict) and isinstance(s.get(k), dict):
                s[k].update(v)
            else:
                s[k] = v
    return s


def save_settings(changes):
    """Nur die genannten Felder in website.json aendern (Rest und Reihenfolge bleiben)."""
    data = {}
    if os.path.exists(SETTINGS):
        with open(SETTINGS, encoding="utf-8-sig") as fh:
            data = json.load(fh)
    for k, v in changes.items():
        if isinstance(v, dict) and isinstance(data.get(k), dict):
            data[k].update(v)
        else:
            data[k] = v
    with open(SETTINGS, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
        fh.write("\n")


def partner_tag(settings=None):
    s = settings if settings is not None else load_settings()
    return (os.environ.get("PARTNER_TAG") or s.get("partner_tag") or "").strip()


def site_url(settings=None):
    """Adresse der Website ohne / am Ende, z. B. https://keroxzz.github.io/dealbar (leer, wenn unbekannt)."""
    s = settings if settings is not None else load_settings()
    dom = (s.get("domain") or "").strip().rstrip("/")
    if dom:
        return dom if dom.startswith("http") else "https://" + dom
    repo = os.environ.get("GITHUB_REPOSITORY", "")
    if "/" in repo:
        owner, name = repo.split("/", 1)
        if name.lower() == f"{owner.lower()}.github.io":
            return f"https://{owner.lower()}.github.io"
        return f"https://{owner.lower()}.github.io/{name}"
    return ""


def impressum_ok(settings=None):
    s = settings if settings is not None else load_settings()
    i = s.get("impressum", {})
    return all(str(i.get(k, "")).strip() for k in ("name", "strasse", "plz_ort", "email")) and \
        not PLACEHOLDER.search(json.dumps(i, ensure_ascii=False))


def channel_link(settings=None):
    s = settings if settings is not None else load_settings()
    k = (s.get("telegram_kanal") or "").strip()
    return f"https://t.me/{k[1:]}" if k.startswith("@") else ""


# --------------------------------------------------------------------------- Tabelle

_held = 0


def _acquire(timeout):
    start = time.time()
    while True:
        try:
            fd = os.open(LOCK, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, str(os.getpid()).encode())
            os.close(fd)
            return
        except FileExistsError:
            try:  # verwaiste Sperre (Programm abgestuerzt) nach 10 Minuten aufheben
                if time.time() - os.path.getmtime(LOCK) > 600:
                    os.remove(LOCK)
                    continue
            except OSError:
                pass
            if time.time() - start > timeout:
                raise TimeoutError("deals.xlsx ist gerade gesperrt, bitte gleich nochmal versuchen.")
            time.sleep(0.2)


@contextlib.contextmanager
def locked(timeout=300):
    """Sperre fuer deals.xlsx (auch fuer Git-Befehle, die die Datei austauschen). Mehrfach verschachtelbar."""
    global _held
    if not _held:
        _acquire(timeout)
    _held += 1
    try:
        yield
    finally:
        _held -= 1
        if not _held:
            try:
                os.remove(LOCK)
            except OSError:
                pass


def save_atomic(wb, path=TABLE):
    tmp = path + ".tmp"
    wb.save(tmp)
    try:
        os.replace(tmp, path)
    except PermissionError:
        os.remove(tmp)
        raise PermissionError("deals.xlsx ist noch in Excel geöffnet. Bitte Excel schließen und erneut starten.")


class Table:
    def __init__(self, wb):
        self.wb = wb
        self.ws = wb["Deals"]
        self.head = {str(c.value).strip(): c.column for c in self.ws[1] if c.value}

    def get(self, r, name):
        col = self.head.get(name)
        return self.ws.cell(r, col).value if col else None

    def text(self, r, name):
        v = self.get(r, name)
        return "" if v is None else str(v).strip()

    def set(self, r, name, value):
        col = self.head.get(name)
        if col:
            self.ws.cell(r, col).value = value

    def nr(self, r):
        try:
            return int(float(self.get(r, "Nr")))
        except (TypeError, ValueError):
            return None

    def rows(self):
        """Zeilennummern aller Deals (Zeilen mit Produktname), in Tabellen-Reihenfolge."""
        return [r for r in range(2, self.ws.max_row + 1) if self.text(r, "Produktname")]

    def row(self, r):
        d = {name: self.get(r, name) for name in self.head}
        d["_row"] = r
        return d

    def find(self, nr=None, folder=None, title=None):
        for r in self.rows():
            if nr is not None and self.nr(r) == int(nr):
                return r
        for r in self.rows():
            if folder and self.text(r, "Ordner") == folder:
                return r
        for r in self.rows():
            if title and self.text(r, "Produktname") == title:
                return r
        return None

    def next_row(self):
        return next((r for r in range(2, self.ws.max_row + 2) if not self.text(r, "Produktname")),
                    self.ws.max_row + 1)

    def next_nr(self):
        return max((self.nr(r) or 0 for r in self.rows()), default=0) + 1

    def ensure_formula(self, r):
        if "Rabatt" in self.head and not self.get(r, "Rabatt"):
            o = self.ws.cell(r, self.head["Alter Preis (€)"]).coordinate
            n = self.ws.cell(r, self.head["Neuer Preis (€)"]).coordinate
            self.set(r, "Rabatt", f'=IF(AND(ISNUMBER({o}),ISNUMBER({n}),{o}>{n}),1-{n}/{o},"")')

    def migrate(self):
        """Fehlende Spalten ergaenzen, Deal-Nummern und ASIN nachtragen. True, wenn sich etwas geaendert hat."""
        changed = False
        ws = self.ws
        for name, width in NEW_COLUMNS:
            if name in self.head:
                continue
            col = max(self.head.values()) + 1
            ws.cell(1, col).value = name
            ws.cell(1, col)._style = copy.copy(ws.cell(1, col - 1)._style)
            proto = self.head.get("Videodatei", col - 1)
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, proto).has_style:
                    ws.cell(r, col)._style = copy.copy(ws.cell(r, proto)._style)
            ws.column_dimensions[get_column_letter(col)].width = width
            self.head[name] = col
            changed = True
        nxt = self.next_nr()
        for r in self.rows():  # Nummern nur fuer echte Deals (Beispielzeilen ohne Link bekommen keine)
            if not self.nr(r) and valid_link(self.text(r, "Affiliate-Link")):
                self.set(r, "Nr", nxt)
                nxt += 1
                changed = True
            if not self.text(r, "ASIN"):
                a = asin_of(self.text(r, "Affiliate-Link"))
                if a:
                    self.set(r, "ASIN", a)
                    changed = True
        lists = {"Kategorie": '"' + ",".join(CATEGORIES) + '"', "Status": f'"{STATUS_LIST}"'}
        for dv in ws.data_validations.dataValidation:
            for name, formula in lists.items():
                col = get_column_letter(self.head.get(name, 0) or 1)
                if name in self.head and str(dv.sqref).startswith(col + "2") and dv.formula1 != formula:
                    dv.formula1 = formula
                    changed = True
        return changed


@contextlib.contextmanager
def open_table(write=False):
    """with open_table(write=True) as t: ...  -> speichert am Ende (nur ohne Fehler)."""
    if not write:
        yield Table(load_workbook(TABLE))
        return
    with locked():
        t = Table(load_workbook(TABLE))
        t.migrate()
        yield t
        save_atomic(t.wb)


def migrate_table():
    """Spalten Nr/ASIN ergaenzen, falls noetig. True, wenn gespeichert wurde."""
    with locked():
        t = Table(load_workbook(TABLE))
        if t.migrate():
            save_atomic(t.wb)
            return True
    return False


# --------------------------------------------------------------------------- Konflikte zusammenfuehren

def _same(t1, r1, t2, r2):
    """Derselbe Deal? Gleicher Ordner, sonst gleicher Link, sonst gleicher Titel."""
    for c in ("Ordner", "Affiliate-Link", "Produktname"):
        a, b = t1.text(r1, c), t2.text(r2, c)
        if a or b:
            return a == b
    return True


def _match(t, src, r):
    """Zeile r der Tabelle src in Tabelle t finden: gleiche Nummer und derselbe Deal, sonst gleicher Ordner
    (die Nummer kann sich bei einem frueheren Zusammenfuehren geaendert haben)."""
    nr, folder, title = src.nr(r), src.text(r, "Ordner"), src.text(r, "Produktname")
    rows = t.rows()
    hit = next((x for x in rows if nr and t.nr(x) == nr and _same(t, x, src, r)), None)
    if hit is None and folder:
        hit = next((x for x in rows if t.text(x, "Ordner") == folder), None)
    if hit is None and not folder:  # Zeile ohne Ordner: gleicher Link und gleicher Titel
        hit = next((x for x in rows if not t.text(x, "Ordner") and t.text(x, "Produktname") == title
                    and _same(t, x, src, r)), None)
    return hit


def merge_tables(base, ours, theirs):
    """Drei Versionen von deals.xlsx (bytes) zeilenweise zusammenfuehren -> bytes.

    ours = Stand, auf den aufgesetzt wird; theirs = die eigene Aenderung, die gewinnt, wenn
    beide Seiten dieselbe Zelle geaendert haben. Zeilen werden nie geloescht.
    """
    tabs = []
    for data in (base, ours, theirs):
        t = Table(load_workbook(io.BytesIO(data))) if data else None
        if t:
            t.migrate()
        tabs.append(t)
    tb, to, tt = tabs
    cols = [c for c in tt.head if c in to.head and c != "Rabatt"]
    for r in tt.rows():
        mine = tt.row(r)
        rb = _match(tb, tt, r) if tb else None
        ro = _match(to, tt, r)
        if ro:
            base_row = tb.row(rb) if rb else {}
            for c in cols:
                o, a, b = base_row.get(c), to.get(ro, c), mine.get(c)
                if a != b and (a == o or b != o):
                    to.set(ro, c, b)
            continue
        if rb:  # auf der anderen Seite geloescht: so lassen
            continue
        nr = tt.nr(r)
        taken = nr and any(to.nr(x) == nr for x in to.rows())
        ro = to.next_row()
        for c in cols:
            to.set(ro, c, mine.get(c))
        if taken:  # beide Seiten haben unabhaengig einen neuen Deal mit derselben Nummer angelegt
            to.set(ro, "Nr", to.next_nr())
        to.ensure_formula(ro)
    out = io.BytesIO()
    to.wb.save(out)
    return out.getvalue()
