"""
dealbar: baut aus deals.xlsx eine statische Website im Ordner "website".
Wird von WEBSITE.bat aufgerufen. Einstellungen und Impressum-Daten stehen in website.json.

Regeln (Amazon-Partnerprogramm):
- Euro-Preise werden nur gezeigt, wenn "preise_in_euro_zeigen" an ist (erst mit Creators API sinnvoll).
- Der Rabatt erscheint nur, solange "Preis geprüft am" juenger als "rabatt_zeigen_stunden" ist,
  immer mit Datum und Uhrzeit. Danach steht nur noch "Aktuellen Preis bei Amazon ansehen".
- Bilder: zuerst SiteStripe-Bild (Datei sitestripe.txt im Deal-Ordner), sonst eigene Fotos aus dem Ordner.
  Screenshots von Amazon werden nicht verwendet.
"""

import html
import json
import os
import re
import shutil
import sys
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402

import brand as brandlib  # noqa: E402

TABLE = os.path.join(ROOT, "deals.xlsx")
DEALS_DIR = os.path.join(ROOT, "deals")
OUT = os.path.join(ROOT, "website")
SETTINGS = os.path.join(ROOT, "website.json")
IMG_EXT = (".jpg", ".jpeg", ".png", ".webp")
HIDDEN = ("abgelaufen", "aus", "offline", "fehler")
PLACEHOLDER = re.compile(r"VORNAME|NACHNAME|STRASSE|PLZ ORT|BEISPIEL\.DE")
DISCLOSURE = "Als Amazon-Partner verdiene ich an qualifizierten Verkäufen."


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("€", "").strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_time(v):
    """'17.10.2026, 15:30', '17.10.2026 15:30 Uhr' oder datetime -> ISO-Text (lokale Zeit)."""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%dT%H:%M")
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})\D*(\d{1,2})?:?(\d{2})?", str(v or ""))
    if not m:
        return None
    d, mo, y, h, mi = m.groups()
    return f"{int(y):04d}-{int(mo):02d}-{int(d):02d}T{int(h or 0):02d}:{int(mi or 0):02d}"


def slug(s):
    s = s.lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:40] or "deal"


def sitestripe_image(folder):
    """Bild-URL aus dem SiteStripe-Code (Text > Bild) in deals/<ordner>/sitestripe.txt."""
    p = os.path.join(DEALS_DIR, folder, "sitestripe.txt")
    if not os.path.exists(p):
        return None
    with open(p, encoding="utf-8", errors="ignore") as fh:
        code = fh.read()
    m = re.search(r"<img[^>]+src=[\"']([^\"']+)", code) or re.search(r"(https?://\S+)", code)
    if not m:
        return None
    src = html.unescape(m.group(1))
    return "https:" + src if src.startswith("//") else src


def own_image(folder, name):
    """Erstes eigenes Foto aus dem Deal-Ordner, verkleinert als webp. Screenshots werden uebersprungen."""
    fpath = os.path.join(DEALS_DIR, folder)
    if not os.path.isdir(fpath):
        return None, f"Ordner nicht gefunden: deals\\{folder}"
    imgs = sorted(f for f in os.listdir(fpath) if f.lower().endswith(IMG_EXT))
    shots = [f for f in imgs if f.lower().startswith(("screenshot", "bildschirmfoto"))]
    imgs = [f for f in imgs if f not in shots]
    if not imgs:
        if shots:
            return None, "nur Screenshots im Ordner (nicht erlaubt), bitte sitestripe.txt oder eigenes Foto"
        return None, "kein Bild (eigenes Foto oder sitestripe.txt)"
    im = Image.open(os.path.join(fpath, imgs[0]))
    im = im.convert("RGBA") if im.mode in ("P", "LA") else im
    bg = Image.new("RGB", im.size, (255, 255, 255))
    bg.paste(im, mask=im.split()[3] if im.mode == "RGBA" else None)
    bg.thumbnail((640, 640), Image.LANCZOS)
    rel = f"bilder/{name}.webp"
    bg.save(os.path.join(OUT, rel), "WEBP", quality=82)
    return rel, None


def read_deals(settings):
    wb = load_workbook(TABLE, data_only=False)
    ws = wb["Deals"]
    head = {str(c.value).strip(): c.column for c in ws[1] if c.value}

    def val(r, name):
        return ws.cell(row=r, column=head[name]).value if name in head else None

    deals, warnings = [], []
    for r in range(2, ws.max_row + 1):
        title = str(val(r, "Produktname") or "").strip()
        if not title:
            continue
        status = str(val(r, "Status") or "").strip().lower()
        if status in HIDDEN:
            continue
        link = str(val(r, "Affiliate-Link") or "").strip()
        if not re.match(r"https?://", link) or "DEIN-LINK" in link:
            warnings.append(f"Zeile {r} ({title[:40]}): kein gültiger Affiliate-Link, nicht auf der Website.")
            continue
        folder = str(val(r, "Ordner") or "").strip()
        name = slug(folder or title)
        img = sitestripe_image(folder) if folder else None
        if not img:
            img, why = own_image(folder, name) if folder else (None, "Spalte 'Ordner' ist leer")
            if why:
                warnings.append(f"Zeile {r} ({title[:40]}): {why}. Wird ohne Bild gezeigt.")

        p_old, p_new = num(val(r, "Alter Preis (€)")), num(val(r, "Neuer Preis (€)"))
        discount = round((p_old - p_new) / p_old * 100) if (p_old and p_new and p_old > p_new) else None
        checked = parse_time(val(r, "Preis geprüft am"))
        show_price = str(val(r, "Preis zeigen") or "ja").strip().lower() not in ("nein", "n", "no")
        euro = settings.get("preise_in_euro_zeigen", False) and show_price
        deals.append({
            "id": name,
            "title": title,
            "hook": str(val(r, "Hook") or "").strip(),
            "label": str(val(r, "Label") or "").strip(),
            "cat": str(val(r, "Kategorie") or "").strip(),
            "link": link,
            "img": img,
            "discount": discount if show_price else None,
            "old": p_old if euro and discount else None,
            "new": p_new if euro else None,
            "checked": checked,
            "added": parse_time(val(r, "Erstellt am")) or checked or "",
            "row": r,
        })
    # Neueste zuerst; bei gleichem Datum gewinnt die spaetere Zeile
    deals.sort(key=lambda d: (d["checked"] or d["added"], d["row"]), reverse=True)
    for d in deals:
        del d["row"]
    return deals, warnings


# --------------------------------------------------------------------------- HTML

def css(brand):
    return """
@font-face{font-family:Inter;font-weight:500;font-display:swap;src:url(fonts/Inter-Medium.otf)}
@font-face{font-family:Inter;font-weight:700;font-display:swap;src:url(fonts/Inter-Bold.otf)}
@font-face{font-family:Inter;font-weight:900;font-display:swap;src:url(fonts/Inter-Black.otf)}
:root{--ink:%(ink)s;--paper:%(paper)s;--bg:%(background)s;--accent:%(accent)s;--muted:%(muted)s;
--line:rgba(17,17,17,.1);--r:14px}
*{box-sizing:border-box}
html{-webkit-text-size-adjust:100%%}
body{margin:0;background:var(--bg);color:var(--ink);font:500 16px/1.5 Inter,system-ui,-apple-system,"Segoe UI",sans-serif}
a{color:inherit}
.wrap{max-width:1120px;margin:0 auto;padding:0 16px}
.ad-note{background:var(--ink);color:var(--paper);font-size:13px;text-align:center;padding:7px 16px}
header.top{padding:28px 0 8px}
.brand{display:flex;align-items:center;gap:12px;text-decoration:none}
.brand img{height:44px;width:auto}
.tagline{color:var(--muted);font-size:12px;letter-spacing:.14em;text-transform:uppercase;margin:6px 0 0}
.controls{position:sticky;top:0;z-index:5;background:var(--bg);padding:14px 0 12px;border-bottom:1px solid var(--line)}
.search{width:100%%;font:inherit;padding:11px 14px;border:1px solid var(--line);border-radius:10px;background:var(--paper);color:var(--ink)}
.search:focus,.chip:focus-visible,.btn:focus-visible,select:focus{outline:2px solid var(--accent);outline-offset:2px}
.row{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin-top:10px}
.row .lbl{font-size:13px;color:var(--muted);margin-right:2px}
.chip{font:inherit;font-size:14px;padding:6px 13px;border-radius:999px;border:1px solid var(--line);background:var(--paper);color:var(--ink);cursor:pointer}
.chip[aria-pressed=true]{background:var(--ink);color:var(--paper);border-color:var(--ink)}
select{font:inherit;font-size:14px;padding:6px 10px;border-radius:10px;border:1px solid var(--line);background:var(--paper);color:var(--ink)}
.bar{display:flex;align-items:center;justify-content:space-between;margin-top:14px}
.count{font-size:13px;color:var(--muted);margin:0}
.legal .box>h2:first-child{margin-top:0}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(250px,1fr));gap:16px;padding:14px 0 40px}
.card{background:var(--paper);border-radius:var(--r);overflow:hidden;display:flex;flex-direction:column;border:1px solid var(--line);scroll-margin-top:150px}
.card:target{outline:3px solid var(--accent)}
.pic{position:relative;aspect-ratio:1/1;background:#fff;display:flex;align-items:center;justify-content:center}
.pic img{max-width:86%%;max-height:86%%;object-fit:contain}
.pic .noimg{color:var(--muted);font-size:13px}
.pic img+.fb{display:none}
.badge{position:absolute;top:10px;left:10px;font-size:11px;font-weight:700;letter-spacing:.08em;background:var(--ink);color:var(--paper);padding:4px 8px;border-radius:6px}
.off{position:absolute;top:10px;right:10px;font-weight:900;font-size:17px;background:var(--accent);color:#fff;padding:5px 9px;border-radius:8px}
.body{padding:14px 16px 16px;display:flex;flex-direction:column;gap:6px;flex:1}
.cat{font-size:12px;color:var(--muted);text-transform:uppercase;letter-spacing:.1em}
.card h2{font-size:16px;line-height:1.35;margin:0;font-weight:700}
.hook{font-size:14px;color:#444;margin:0}
.price{display:flex;gap:8px;align-items:baseline;margin-top:auto;padding-top:6px}
.price .new{font-weight:900;font-size:22px}
.price .old{color:var(--muted);text-decoration:line-through;font-size:14px}
.stamp{font-size:12px;color:var(--muted);margin:0}
.btn{display:block;text-align:center;margin-top:10px;padding:11px 14px;border-radius:10px;background:var(--accent);color:#fff;font-weight:700;text-decoration:none}
.btn:hover{filter:brightness(.94)}
.empty{text-align:center;color:var(--muted);padding:60px 0}
footer{border-top:1px solid var(--line);padding:24px 0 40px;font-size:13px;color:var(--muted)}
footer nav{display:flex;gap:18px;flex-wrap:wrap;margin-bottom:10px}
footer nav a{color:var(--ink);font-weight:700;text-decoration:none}
.legal{max-width:760px;padding:24px 16px 48px;margin:0 auto}
.legal h1{font-size:30px;font-weight:900;margin:8px 0 18px}
.legal h2{font-size:18px;margin:26px 0 6px}
.legal .box{background:var(--paper);border:1px solid var(--line);border-radius:var(--r);padding:20px 22px}
.warn{background:#fff3cd;border:1px solid #e0c36a;color:#5c4500;padding:10px 14px;border-radius:10px;font-size:14px}
@media (max-width:560px){.controls{position:static}.row{flex-wrap:nowrap;overflow-x:auto;scrollbar-width:none;margin-right:-16px;padding-right:16px}
.chip{flex:none}.grid{grid-template-columns:1fr 1fr;gap:10px}
.card h2{font-size:14px}.hook{display:none}.body{padding:10px 12px 12px}.off{font-size:14px}.btn{padding:10px 8px;font-size:14px}}
@media (max-width:360px){.grid{grid-template-columns:1fr}}
""" % brand


def page(settings, brand, title, main, extra_head="", script=""):
    desc = html.escape(settings.get("beschreibung", ""))
    domain = settings.get("domain", "").rstrip("/")
    og = f'<meta property="og:url" content="{html.escape(domain)}/">' if domain else ""
    return f"""<!doctype html>
<html lang="de">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title>
<meta name="description" content="{desc}">
<meta property="og:title" content="{html.escape(title)}">
<meta property="og:description" content="{desc}">
<meta property="og:image" content="logo.png">
{og}
<link rel="icon" href="favicon.png">
<style>{css(brand)}</style>
{extra_head}
</head>
<body>
<div class="ad-note">Anzeige: Links mit * sind Amazon-Partnerlinks. {DISCLOSURE}</div>
<header class="top"><div class="wrap">
<a class="brand" href="index.html" aria-label="{html.escape(brand['name_bold'] + brand['name_light'])} Startseite"><img src="logo.png" alt="{html.escape(brand['name_bold'] + brand['name_light'])}"></a>
<p class="tagline">{html.escape(brand.get('tagline', ''))}</p>
</div></header>
{main}
<footer><div class="wrap">
<nav><a href="index.html">Deals</a><a href="affiliate-hinweis.html">Affiliate-Hinweis</a><a href="impressum.html">Impressum</a><a href="datenschutz.html">Datenschutz</a></nav>
<p>* Amazon-Partnerlink (Werbung). {DISCLOSURE} Für dich ändert sich der Preis dadurch nicht.<br>
Rabatte gelten zum angegebenen Zeitpunkt und können sich jederzeit ändern. Maßgeblich ist der Preis auf Amazon zum Zeitpunkt des Kaufs.
Amazon und das Amazon-Logo sind Warenzeichen von Amazon.com, Inc. oder eines seiner verbundenen Unternehmen.</p>
</div></footer>
{script}
</body>
</html>
"""


INDEX_MAIN = """
<main>
<div class="controls"><div class="wrap">
<input class="search" id="q" type="search" placeholder="Deals durchsuchen …" aria-label="Deals durchsuchen">
<div class="row" id="cats" hidden><span class="lbl">Kategorie</span></div>
<div class="row"><span class="lbl">Rabatt</span>
<button class="chip" data-min="0" aria-pressed="true">Alle</button>
<button class="chip" data-min="20" aria-pressed="false">ab 20 %</button>
<button class="chip" data-min="30" aria-pressed="false">ab 30 %</button>
<button class="chip" data-min="50" aria-pressed="false">ab 50 %</button>
</div>
</div></div>
<div class="wrap">
<div class="bar"><p class="count" id="count"></p>
<select id="sort" aria-label="Sortierung"><option value="new">Neueste zuerst</option><option value="off">Höchster Rabatt</option></select></div>
<div class="grid" id="grid"></div>
<p class="empty" id="empty" hidden>Keine Deals für diese Auswahl. Probier einen anderen Filter.</p>
</div>
</main>
"""

INDEX_JS = """
<script id="data" type="application/json">%(data)s</script>
<script>
(function(){
var deals=JSON.parse(document.getElementById('data').textContent);
var MAX_H=%(maxh)s, now=Date.now();
var state={cat:'',min:0,q:'',sort:'new'};
function esc(s){return String(s).replace(/[&<>"']/g,function(c){return {'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]})}
function euro(v){return v.toLocaleString('de-DE',{style:'currency',currency:'EUR'})}
function stamp(iso){var d=new Date(iso);return d.toLocaleDateString('de-DE',{day:'2-digit',month:'2-digit',year:'numeric'})+', '+d.toLocaleTimeString('de-DE',{hour:'2-digit',minute:'2-digit'})+' Uhr'}
deals.forEach(function(d){
  d.fresh=!!d.checked && (now-new Date(d.checked).getTime())/36e5 < MAX_H;
  d.off=d.fresh&&d.discount?d.discount:0;
});
var cats=[];deals.forEach(function(d){if(d.cat&&cats.indexOf(d.cat)<0)cats.push(d.cat)});cats.sort();
var catRow=document.getElementById('cats');
if(cats.length>1){catRow.hidden=false;[''].concat(cats).forEach(function(c){
  var b=document.createElement('button');b.className='chip';b.textContent=c||'Alle';b.dataset.cat=c;
  b.setAttribute('aria-pressed',c===''?'true':'false');catRow.appendChild(b)})}
function press(row,btn){row.querySelectorAll('.chip').forEach(function(x){x.setAttribute('aria-pressed',x===btn?'true':'false')})}
catRow.addEventListener('click',function(e){var b=e.target.closest('.chip');if(!b)return;state.cat=b.dataset.cat;press(catRow,b);render()});
document.querySelectorAll('[data-min]').forEach(function(b){b.addEventListener('click',function(){state.min=+b.dataset.min;press(b.parentNode,b);render()})});
document.getElementById('q').addEventListener('input',function(e){state.q=e.target.value.trim().toLowerCase();render()});
document.getElementById('sort').addEventListener('change',function(e){state.sort=e.target.value;render()});
function card(d){
  var h='<article class="card" id="'+esc(d.id)+'"><div class="pic">';
  h+=d.img?'<img src="'+esc(d.img)+'" alt="'+esc(d.title)+'" loading="lazy" referrerpolicy="no-referrer" onerror="this.remove()"><span class="noimg fb">Bild bei Amazon</span>':'<span class="noimg">Bild bei Amazon</span>';
  if(d.label)h+='<span class="badge">'+esc(d.label)+'</span>';
  if(d.off)h+='<span class="off">−'+d.off+' %%</span>';
  h+='</div><div class="body">';
  if(d.cat)h+='<span class="cat">'+esc(d.cat)+'</span>';
  h+='<h2>'+esc(d.title)+'</h2>';
  if(d.hook)h+='<p class="hook">'+esc(d.hook)+'</p>';
  if(d.fresh&&d.new){h+='<div class="price"><span class="new">'+euro(d.new)+'*</span>'+(d.old?'<span class="old">'+euro(d.old)+'</span>':'')+'</div>'}
  else h+='<div class="price"></div>';
  if(d.fresh&&(d.off||d.new))h+='<p class="stamp">'+(d.new?'Preis':'Rabatt')+' vom '+stamp(d.checked)+', kann sich jederzeit ändern.</p>';
  else h+='<p class="stamp">Aktuellen Preis bei Amazon ansehen.</p>';
  h+='<a class="btn" href="'+esc(d.link)+'" target="_blank" rel="sponsored nofollow noopener">Zum Deal bei Amazon*</a>';
  return h+'</div></article>';
}
function render(){
  var list=deals.filter(function(d){
    if(state.cat&&d.cat!==state.cat)return false;
    if(state.min&&d.off<state.min)return false;
    if(state.q&&(d.title+' '+d.hook+' '+d.cat+' '+d.label).toLowerCase().indexOf(state.q)<0)return false;
    return true});
  if(state.sort==='off')list=list.slice().sort(function(a,b){return b.off-a.off});
  document.getElementById('grid').innerHTML=list.map(card).join('');
  document.getElementById('empty').hidden=list.length>0;
  document.getElementById('count').textContent=list.length+(list.length===1?' Deal':' Deals');
  if(location.hash){var t=document.getElementById(location.hash.slice(1));if(t&&!render.done){t.scrollIntoView();render.done=true}}
}
render();
})();
</script>
"""


def legal_main(title, body):
    return f'<main class="legal"><h1>{html.escape(title)}</h1>{body}</main>'


def impressum_body(settings):
    i = {k: html.escape(v) for k, v in settings.get("impressum", {}).items()}
    tel = f"<br>Telefon: {i['telefon']}" if i.get("telefon") else ""
    warn = ""
    if PLACEHOLDER.search(json.dumps(settings.get("impressum", {}))):
        warn = '<p class="warn">Bitte in website.json deine echten Angaben eintragen und die Website neu bauen.</p>'
    return f"""{warn}<div class="box">
<h2>Angaben gemäß § 5 DDG</h2>
<p>{i.get('name', '')}<br>{i.get('strasse', '')}<br>{i.get('plz_ort', '')}</p>
<h2>Kontakt</h2>
<p>E-Mail: <a href="mailto:{i.get('email', '')}">{i.get('email', '')}</a>{tel}</p>
<h2>Verantwortlich für den Inhalt nach § 18 Abs. 2 MStV</h2>
<p>{i.get('name', '')}, Anschrift wie oben</p>
<h2>Hinweis zu Partnerlinks</h2>
<p>Diese Website enthält Amazon-Partnerlinks. {DISCLOSURE} Mehr dazu im <a href="affiliate-hinweis.html">Affiliate-Hinweis</a>.</p>
<h2>Haftung für Links</h2>
<p>Diese Website enthält Links zu externen Websites, auf deren Inhalte ich keinen Einfluss habe. Für diese Inhalte ist der jeweilige Anbieter verantwortlich. Bei Bekanntwerden von Rechtsverletzungen entferne ich solche Links umgehend.</p>
<h2>Verbraucherstreitbeilegung</h2>
<p>Ich bin nicht bereit und nicht verpflichtet, an Streitbeilegungsverfahren vor einer Verbraucherschlichtungsstelle teilzunehmen.</p>
</div>"""


def affiliate_body():
    return f"""<div class="box">
<p><strong>{DISCLOSURE}</strong></p>
<p>dealbar ist Teilnehmer des Partnerprogramms von Amazon EU, das zur Bereitstellung eines Mediums für Websites konzipiert wurde, mittels dessen durch die Platzierung von Werbeanzeigen und Links zu Amazon.de Werbekostenerstattung verdient werden kann.</p>
<p>Alle Links zu Amazon auf dieser Seite sind Partnerlinks und mit einem Sternchen (*) gekennzeichnet. Kaufst du über einen solchen Link, erhalte ich eine kleine Provision. Für dich ändert sich der Preis dadurch nicht.</p>
<h2>Preise und Rabatte</h2>
<p>Rabatte werden mit Datum und Uhrzeit der Prüfung angezeigt und nur für kurze Zeit. Preise und Verfügbarkeit können sich jederzeit ändern. Maßgeblich ist allein der Preis, der zum Zeitpunkt des Kaufs auf Amazon.de angezeigt wird.</p>
<h2>Auswahl der Deals</h2>
<p>Ich wähle die Angebote selbst aus. Die Provision hat keinen Einfluss darauf, welche Produkte ich zeige.</p>
</div>"""


def datenschutz_body(settings, uses_amazon_images):
    i = {k: html.escape(v) for k, v in settings.get("impressum", {}).items()}
    img_txt = ""
    if uses_amazon_images:
        img_txt = """<h2>Produktbilder von Amazon</h2>
<p>Einige Produktbilder werden direkt von Servern von Amazon geladen (Amazon-Partnerprogramm, SiteStripe). Dabei wird deine IP-Adresse an Amazon übertragen. Rechtsgrundlage ist mein berechtigtes Interesse an einer korrekten Darstellung der Angebote (Art. 6 Abs. 1 lit. f DSGVO). Anbieter: Amazon Europe Core S.à r.l., 38 Avenue John F. Kennedy, L-1855 Luxemburg.</p>"""
    return f"""<div class="box">
<h2>Verantwortlicher</h2>
<p>{i.get('name', '')}, {i.get('strasse', '')}, {i.get('plz_ort', '')}<br>E-Mail: {i.get('email', '')}</p>
<h2>Keine Cookies, kein Tracking</h2>
<p>Diese Website setzt keine Cookies und verwendet keine Analyse- oder Tracking-Werkzeuge. Schriftarten werden lokal ausgeliefert, nicht von Google.</p>
<h2>Server-Logfiles</h2>
<p>Beim Aufruf der Seite speichert der Hosting-Anbieter technisch notwendige Daten (z. B. IP-Adresse, Datum und Uhrzeit, aufgerufene Seite, Browser). Das dient dem sicheren Betrieb der Website (Art. 6 Abs. 1 lit. f DSGVO). Die Daten werden nicht mit anderen Daten zusammengeführt und nach kurzer Zeit gelöscht.</p>
{img_txt}
<h2>Amazon-Partnerlinks</h2>
<p>Klickst du auf einen Partnerlink, wirst du zu Amazon weitergeleitet. Amazon kann dort Cookies setzen, um nachzuvollziehen, dass der Kauf über diese Seite kam. Dafür gilt die Datenschutzerklärung von Amazon. Ich erhalte von Amazon keine personenbezogenen Daten über dich, nur anonyme Umsatzzahlen.</p>
<h2>Deine Rechte</h2>
<p>Du hast das Recht auf Auskunft, Berichtigung, Löschung, Einschränkung der Verarbeitung, Datenübertragbarkeit und Widerspruch (Art. 15 bis 21 DSGVO) sowie das Recht, dich bei einer Datenschutz-Aufsichtsbehörde zu beschweren.</p>
</div>"""


def build():
    print("dealbar Website-Generator\n")
    if not os.path.exists(TABLE):
        sys.exit(f"Tabelle nicht gefunden: {TABLE}")
    with open(SETTINGS, encoding="utf-8") as fh:
        settings = json.load(fh)
    brand = brandlib.load_brand(os.path.join(ROOT, "marke.json"))

    if os.path.isdir(OUT):
        shutil.rmtree(OUT)
    os.makedirs(os.path.join(OUT, "bilder"))
    os.makedirs(os.path.join(OUT, "fonts"))
    for f in ("Inter-Medium.otf", "Inter-Bold.otf", "Inter-Black.otf", "LICENSE.txt"):
        src = os.path.join(HERE, "fonts", f)
        if os.path.exists(src):
            shutil.copy(src, os.path.join(OUT, "fonts", f))

    # Logo und Favicon aus der Marke
    from deal_video import font
    ink = brandlib.hex_rgb(brand["ink"])
    brandlib.make_lockup(brand, 120, font).save(os.path.join(OUT, "logo.png"))
    mark = brandlib.make_mark(96, ink)
    fav = Image.new("RGBA", (128, 128), (0, 0, 0, 0))
    fav.alpha_composite(mark, ((128 - mark.width) // 2, (128 - mark.height) // 2))
    fav.save(os.path.join(OUT, "favicon.png"))

    deals, warnings = read_deals(settings)
    data = json.dumps(deals, ensure_ascii=False).replace("</", "<\\/")
    script = INDEX_JS % {"data": data, "maxh": float(settings.get("rabatt_zeigen_stunden", 24))}
    site = settings.get("titel", "dealbar")
    pages = {
        "index.html": page(settings, brand, site, INDEX_MAIN, script=script),
        "impressum.html": page(settings, brand, "Impressum – dealbar", legal_main("Impressum", impressum_body(settings))),
        "datenschutz.html": page(settings, brand, "Datenschutz – dealbar", legal_main(
            "Datenschutzerklärung", datenschutz_body(settings, any(d["img"] and d["img"].startswith("http") for d in deals)))),
        "affiliate-hinweis.html": page(settings, brand, "Affiliate-Hinweis – dealbar",
                                       legal_main("Affiliate-Hinweis", affiliate_body())),
    }
    for name, content in pages.items():
        with open(os.path.join(OUT, name), "w", encoding="utf-8") as fh:
            fh.write(content)

    print(f"{len(deals)} Deals auf der Website.")
    if PLACEHOLDER.search(json.dumps(settings.get("impressum", {}))):
        warnings.insert(0, "Impressum: in website.json stehen noch Platzhalter. Vor dem Veröffentlichen ausfüllen!")
    if warnings:
        print("\nHinweise:")
        for w in warnings:
            print("  - " + w)
    print("\nFertig. Zum Ansehen website\\index.html doppelklicken.")


if __name__ == "__main__":
    build()
